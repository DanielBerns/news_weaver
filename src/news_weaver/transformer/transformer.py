import os
import sys
import json
import logging
import time
import mimetypes
from datetime import datetime, timezone
from typing import Dict, Any, List, Tuple
from concurrent.futures import ThreadPoolExecutor, as_completed

import yaml
import httpx
from sqlalchemy import create_engine, Column, Integer, String, DateTime, ForeignKey, Text
from sqlalchemy.orm import sessionmaker, declarative_base, Session

# --- Library Imports for Content Extraction ---
try:
    from bs4 import BeautifulSoup
    from pypdf import PdfReader
    from docx import Document as DocxDocument
    from openpyxl import load_workbook
    from PIL import Image
    import pytesseract
except ImportError as e:
    print(f"Missing dependency: {e}")
    print("Please run: uv pip install beautifulsoup4 pypdf python-docx openpyxl pillow pytesseract")
    sys.exit(1)

# --- Configuration & Logging ---

def load_config(config_path: str = "config.yaml") -> dict:
    if not os.path.exists(config_path):
        print(f"Error: {config_path} not found.")
        sys.exit(1)
    with open(config_path, "r") as f:
        return yaml.safe_load(f)

CONFIG = load_config()

# Logging
logging.basicConfig(
    filename=CONFIG["logging"]["file"],
    level=getattr(logging, CONFIG["logging"]["level"].upper(), logging.INFO),
    format='{"timestamp": "%(asctime)s", "level": "%(levelname)s", "component": "Transformer", "message": "%(message)s"}'
)
logger = logging.getLogger("Transformer")

# --- Database Setup (pipeline.db) ---

DATABASE_URL = CONFIG["database"]["pipeline_db_url"]
engine = create_engine(DATABASE_URL)
SessionLocal = sessionmaker(bind=engine)
Base = declarative_base()

class ScrapedFile(Base):
    __tablename__ = "scraped_files"
    id = Column(Integer, primary_key=True)
    source_id = Column(Integer, nullable=False)
    local_path = Column(String, nullable=False)
    filename = Column(String, nullable=False)
    mimetype = Column(String, nullable=False)
    scraped_at = Column(DateTime)
    status = Column(String, default="SCRAPED")
    retry_count = Column(Integer, default=0)
    notes = Column(Text, nullable=True)

class Source(Base):
    __tablename__ = "sources"
    id = Column(Integer, primary_key=True)
    url = Column(String, unique=True)
    # ... other fields ignored for this script

# --- Extraction Logic ---

def extract_text_from_html(filepath: str) -> Tuple[str, str]:
    """Returns (title, text_content)."""
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f, "html.parser")

    # Remove script and style elements
    for script in soup(["script", "style", "nav", "footer"]):
        script.extract()

    title = soup.title.string if soup.title else "No Title"
    text = soup.get_text(separator="\n")

    # Clean up whitespace
    lines = (line.strip() for line in text.splitlines())
    chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
    clean_text = '\n'.join(chunk for chunk in chunks if chunk)

    return title, clean_text

def extract_text_from_pdf(filepath: str) -> str:
    reader = PdfReader(filepath)
    text = []
    for page in reader.pages:
        text.append(page.extract_text() or "")
    return "\n".join(text)

def extract_text_from_docx(filepath: str) -> str:
    doc = DocxDocument(filepath)
    return "\n".join([para.text for para in doc.paragraphs])

def extract_data_from_xlsx(filepath: str) -> List[Dict[str, Any]]:
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.rows)
    if not rows:
        return []

    headers = [cell.value for cell in rows[0]]
    data = []
    for row in rows[1:]:
        row_data = {}
        for i, cell in enumerate(row):
            if i < len(headers):
                row_data[str(headers[i])] = cell.value
        data.append(row_data)
    return data

def process_image(filepath: str) -> Dict[str, Any]:
    """Performs OCR and metadata extraction."""
    try:
        image = Image.open(filepath)

        # 1. OCR (Text Extraction)
        # Ensure tesseract is installed on the system (apt install tesseract-ocr)
        text = pytesseract.image_to_string(image)

        # 2. Metadata (EXIF)
        exif_data = image._getexif()
        metadata = {}
        if exif_data:
            for tag_id, value in exif_data.items():
                # We interpret just a few common tags to keep JSON clean
                # 36867 = DateTimeOriginal, 271 = Make, 272 = Model
                if tag_id in (36867, 271, 272):
                    metadata[str(tag_id)] = str(value)

        # 3. Object Detection (Placeholder)
        # Real object detection requires heavy libraries like Torch/Tensorflow.
        # We return a stub list as per specs to avoid dependency bloat.
        detected_objects = ["object_detection_model_not_loaded"]

        return {
            "extracted_text": text.strip(),
            "metadata": metadata,
            "detected_objects": detected_objects
        }
    except Exception as e:
        logger.error(f"Image processing failed for {filepath}: {e}")
        raise e

# --- Loader API Client ---

def send_to_loader(endpoint: str, payload: dict) -> bool:
    """Sends processed data to the Loader API with retries."""
    api_cfg = CONFIG["api"]
    url = f"http://{api_cfg['host']}:{api_cfg['port']}/{endpoint}"
    headers = {
        "X-API-Key": api_cfg["secret_key"],
        "Content-Type": "application/json"
    }

    # Simple retry mechanism (3 attempts)
    for attempt in range(3):
        try:
            with httpx.Client(timeout=10.0) as client:
                response = client.post(url, json=payload, headers=headers)
                if response.status_code in (200, 201):
                    return True
                elif response.status_code == 409:
                    logger.info(f"Duplicate entry for source_file_id {payload['source_file_id']}")
                    return True # Treat duplicate as success
                else:
                    logger.warning(f"Loader returned {response.status_code}: {response.text}")
        except httpx.RequestError as e:
            logger.warning(f"Loader API unreachable (Attempt {attempt+1}/3): {e}")
            time.sleep(2)

    return False

# --- Main Processing Workflow ---

def process_file_record(file_record: ScrapedFile, db_session: Session):
    """Router function that dispatches file to correct extractor and sends to API."""

    # Check if file exists on disk
    if not os.path.exists(file_record.local_path):
        logger.error(f"File missing on disk: {file_record.local_path}")
        file_record.status = "TRANSFORM_FAILED"
        file_record.notes = "File not found on disk"
        db_session.commit()
        return

    # Get Source URL (needed for API payload)
    source = db_session.query(Source).filter(Source.id == file_record.source_id).first()
    source_url = source.url if source else "unknown_source"

    mime = file_record.mimetype.lower()
    path = file_record.local_path

    try:
        # --- 1. TRANSFORM ---
        endpoint = ""
        payload = {
            "source_file_id": file_record.id,
            "url": source_url,
            "mimetype": mime
        }

        if "html" in mime:
            title, content = extract_text_from_html(path)
            endpoint = "articles"
            payload.update({"title": title, "content": content, "language": "en"})

        elif "pdf" in mime or "wordprocessing" in mime: # docx is usually 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            if "pdf" in mime:
                content = extract_text_from_pdf(path)
            else:
                content = extract_text_from_docx(path)

            endpoint = "documents"
            payload.update({"filename": file_record.filename, "content": content})

        elif "spreadsheet" in mime or "excel" in mime:
            data = extract_data_from_xlsx(path)
            endpoint = "spreadsheets"
            payload.update({"filename": file_record.filename, "data_json": data})

        elif "image" in mime:
            img_data = process_image(path)
            endpoint = "images"
            payload.update({
                "extracted_text": img_data["extracted_text"],
                "detected_objects": img_data["detected_objects"],
                "image_metadata": img_data["metadata"] # Match pydantic schema field name
            })

        else:
            # Fallback for unknown text types
            if "text" in mime:
                with open(path, "r", errors="ignore") as f:
                    content = f.read()
                endpoint = "documents"
                payload.update({"filename": file_record.filename, "content": content})
            else:
                raise ValueError(f"Unsupported mimetype: {mime}")

        # --- 2. LOAD ---
        logger.info(f"Sending processed data for File {file_record.id} to /{endpoint}")
        success = send_to_loader(endpoint, payload)

        if success:
            file_record.status = "PROCESSED_SUCCESSFULLY"
            file_record.notes = None
        else:
            file_record.status = "LOAD_FAILED"
            file_record.notes = "API unreachable or returned error"

    except Exception as e:
        logger.error(f"Transformation failed for File {file_record.id}: {e}")
        file_record.status = "TRANSFORM_FAILED"
        file_record.notes = str(e)

    finally:
        db_session.commit()

def main():
    session = SessionLocal()
    try:
        # Fetch pending files
        # We also pick up LOAD_FAILED to retry them
        pending_files = session.query(ScrapedFile).filter(
            ScrapedFile.status.in_(["SCRAPED", "LOAD_FAILED"])
        ).limit(50).all() # Process in batches of 50

        if not pending_files:
            return # Nothing to do

        logger.info(f"Found {len(pending_files)} files to process.")

        # Mark as processing to prevent race conditions if multiple transformers ran
        # (Though our Cron is serial, this is good practice)
        for f in pending_files:
            f.status = "PROCESSING"
        session.commit()

        # Execute in parallel threads (IO bound due to API calls, CPU bound due to OCR)
        # Using 4 workers is a safe default for a personal machine
        with ThreadPoolExecutor(max_workers=4) as executor:
            # We must map back to a fresh session per thread or handle session strictly
            # Simplified: We pass the object, but we must be careful with SQLite locks.
            # Safer Approach for SQLite: Run sequentially or use a process pool with separate DB engines.
            # Given SQLite's concurrency limitations, we will run SEQUENTIALLY for safety in this version.
            # To enable concurrency with SQLite, we'd need WAL mode enabled.

            for file_record in pending_files:
                process_file_record(file_record, session)

    except Exception as e:
        logger.critical(f"Critical error in Transformer: {e}")
    finally:
        session.close()

if __name__ == "__main__":
    main()
